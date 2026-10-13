"""CSV and Excel exports with round-trip verification."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from config import (
    BOOL_COLUMNS,
    CSV_ENCODING,
    EXPORT_DIR,
    FILENAME_PREFIX,
    FILENAME_TIME_FORMAT,
    ID_COLUMNS,
    PRICE_COLUMNS,
    STANDARD_COLUMNS,
    TEXT_COLUMNS,
)

UTF8_BOM = b"\xef\xbb\xbf"
EXCEL_SHEET_NAME = "Products"


def _timestamp(now: Optional[datetime] = None) -> str:
    return (now or datetime.now()).strftime(FILENAME_TIME_FORMAT)


def make_filename(now: Optional[datetime] = None) -> str:
    """shopee_products_<YYYYMMDD_HHMMSS>.csv"""
    return f"{FILENAME_PREFIX}_{_timestamp(now)}.csv"


def make_excel_filename(now: Optional[datetime] = None) -> str:
    """shopee_products_<YYYYMMDD_HHMMSS>.xlsx"""
    return f"{FILENAME_PREFIX}_{_timestamp(now)}.xlsx"


def _export_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Return the exact 20-column export schema without mutating the source DataFrame."""
    export_df = df.reindex(columns=STANDARD_COLUMNS).copy()
    # Product/shop IDs should remain text in Excel too, avoiding scientific notation.
    for column in ID_COLUMNS:
        export_df[column] = export_df[column].astype("string").fillna("")
    return export_df


def dataframe_to_csv_bytes(df: pd.DataFrame) -> bytes:
    """All 20 standard columns. Missing values are empty cells; ratings print as 4.8, not 4.80."""
    text = _export_frame(df).to_csv(index=False, float_format="%g")
    return text.encode(CSV_ENCODING)       # 'utf-8-sig' adds the BOM


def dataframe_to_excel_bytes(df: pd.DataFrame) -> bytes:
    """Create an Excel workbook with the 20 standard columns and user-friendly formatting."""
    export_df = _export_frame(df)
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name=EXCEL_SHEET_NAME)
        ws = writer.book[EXCEL_SHEET_NAME]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Bold headers; keep styling intentionally simple for classroom/demo use.
        for cell in ws[1]:
            cell.font = Font(bold=True)

        header_to_index = {cell.value: cell.column for cell in ws[1]}
        for column in ID_COLUMNS:
            column_index = header_to_index[column]
            for cell in ws.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
                cell[0].number_format = "@"

        for column in PRICE_COLUMNS:
            column_index = header_to_index[column]
            for cell in ws.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
                cell[0].number_format = "#,##0"

        if "discount_pct" in header_to_index:
            column_index = header_to_index["discount_pct"]
            for cell in ws.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
                cell[0].number_format = "0.##"

        if "rating" in header_to_index:
            column_index = header_to_index["rating"]
            for cell in ws.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
                cell[0].number_format = "0.##"

        # Readable widths, bounded so a long description/URL doesn't make columns enormous.
        for column_cells in ws.columns:
            letter = column_cells[0].column_letter
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:51])
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 45)

    return output.getvalue()


def verify_csv_bytes(data: bytes, df: pd.DataFrame) -> tuple[bool, str]:
    """Read the CSV back and compare with the DataFrame. Returns (ok, message)."""
    if not data.startswith(UTF8_BOM):
        return False, "CSV does not start with a UTF-8 BOM (utf-8-sig)."
    try:
        back = pd.read_csv(io.BytesIO(data), encoding=CSV_ENCODING, dtype=str,
                           keep_default_na=False)
    except Exception:
        return False, "CSV could not be read back."
    if list(back.columns) != STANDARD_COLUMNS:
        return False, "CSV columns differ from the standard schema."
    if len(back) != len(df):
        return False, "CSV row count differs from the DataFrame."
    # Check every text field so Vietnamese text in keyword/shop/location/category is also protected.
    for column in TEXT_COLUMNS:
        original = df[column].astype("string").fillna("").tolist()
        if back[column].tolist() != original:
            return False, f"Text changed during CSV export in column '{column}'."
    return True, f"CSV verified: {len(back)} rows, {len(back.columns)} columns, UTF-8-SIG; text fields preserved."


def verify_excel_bytes(data: bytes, df: pd.DataFrame) -> tuple[bool, str]:
    """Read the Excel workbook back and verify schema, row count, text and ID preservation."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        if EXCEL_SHEET_NAME not in wb.sheetnames:
            return False, f"Excel sheet '{EXCEL_SHEET_NAME}' is missing."
        ws = wb[EXCEL_SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        return False, "Excel file could not be read back."

    if not rows:
        return False, "Excel file is empty."
    headers = list(rows[0])
    if headers != STANDARD_COLUMNS:
        return False, "Excel columns differ from the standard schema."
    if len(rows) - 1 != len(df):
        return False, "Excel row count differs from the DataFrame."

    # Verify every text field, not only the product name.
    for column in TEXT_COLUMNS:
        idx = headers.index(column)
        expected_text = df[column].astype("string").fillna("").tolist()
        actual_text = ["" if row[idx] is None else str(row[idx]) for row in rows[1:]]
        if actual_text != expected_text:
            return False, f"Text changed during Excel export in column '{column}'."

    for column in ID_COLUMNS:
        idx = headers.index(column)
        expected = df[column].astype("string").fillna("").tolist()
        actual = ["" if row[idx] is None else str(row[idx]) for row in rows[1:]]
        if actual != expected:
            return False, f"Excel ID values changed in '{column}'."

    return True, f"Excel verified: {len(rows) - 1} rows, {len(headers)} columns."


def export_csv(df: pd.DataFrame, directory: Path = EXPORT_DIR,
               now: Optional[datetime] = None) -> Path:
    """Write and verify a CSV file on disk, returning its path."""
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / make_filename(now)
    data = dataframe_to_csv_bytes(df)
    path.write_bytes(data)
    ok, message = verify_csv_bytes(path.read_bytes(), df)
    if not ok:
        raise ValueError(message)
    return path


def export_excel(df: pd.DataFrame, directory: Path = EXPORT_DIR,
                  now: Optional[datetime] = None) -> Path:
    """Write the Excel workbook to disk and verify it."""
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / make_excel_filename(now)
    data = dataframe_to_excel_bytes(df)
    path.write_bytes(data)
    ok, message = verify_excel_bytes(path.read_bytes(), df)
    if not ok:
        raise ValueError(message)
    return path
