"""Separate review CSV/Excel exports with Unicode verification and formula safety."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Sequence

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from config import EXPORT_DIR, FILENAME_TIME_FORMAT, REVIEW_COLUMNS, REVIEW_ID_COLUMNS, REVIEW_TEXT_COLUMNS
from utils.validators import validate_selected_columns

UTF8_BOM = b"\xef\xbb\xbf"
SHEET_NAME = "Reviews"
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _timestamp(now: datetime | None = None) -> str:
    return (now or datetime.now()).strftime(FILENAME_TIME_FORMAT)


def make_review_filename(product_id: str, extension: str, now: datetime | None = None) -> str:
    safe_id = "".join(ch for ch in str(product_id) if ch.isalnum() or ch in "-_-")[:80] or "unknown"
    if extension not in {"csv", "xlsx"}:
        raise ValueError("Review export format must be csv or xlsx.")
    return f"shopee_reviews_{safe_id}_{_timestamp(now)}.{extension}"


def _frame(df: pd.DataFrame, columns: Sequence[str] | None = None) -> pd.DataFrame:
    selected = list(REVIEW_COLUMNS) if columns is None else validate_selected_columns(columns, REVIEW_COLUMNS, "review output")
    return df.reindex(columns=selected).copy()


def _is_formula_like(value: object) -> bool:
    return isinstance(value, str) and value.lstrip().startswith(_FORMULA_PREFIXES)


def _safe_csv_frame(df: pd.DataFrame) -> pd.DataFrame:
    safe = df.copy()
    for column in safe.columns:
        safe[column] = safe[column].map(lambda value: "'" + value if _is_formula_like(value) else value)
    return safe


def dataframe_to_review_csv_bytes(df: pd.DataFrame, columns: Sequence[str] | None = None) -> bytes:
    selected = _frame(df, columns)
    return _safe_csv_frame(selected).to_csv(index=False, lineterminator="\n").encode("utf-8-sig")


def dataframe_to_review_excel_bytes(df: pd.DataFrame, columns: Sequence[str] | None = None) -> bytes:
    selected = _frame(df, columns)
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    headers = list(selected.columns)
    for col, header in enumerate(headers, 1):
        c = ws.cell(1, col, header)
        c.font = Font(bold=True)
        c.data_type = "s"
    for r, record in enumerate(selected.itertuples(index=False, name=None), start=2):
        for c, (column, value) in enumerate(zip(headers, record), start=1):
            if pd.isna(value):
                value = None
            cell = ws.cell(r, c, value)
            if value is not None and (isinstance(value, str) or column in REVIEW_ID_COLUMNS):
                cell.value = str(value)
                # Force an explicit string; review text must never be interpreted as a formula.
                cell.data_type = "s"
            if column in REVIEW_ID_COLUMNS:
                cell.number_format = "@"
    ws.freeze_panes = "A2"
    if headers:
        ws.auto_filter.ref = ws.dimensions
    for cells in ws.columns:
        letter = cells[0].column_letter
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in cells[:51])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 50)
    wb.save(output)
    return output.getvalue()


def verify_review_csv_bytes(data: bytes, df: pd.DataFrame, columns: Sequence[str] | None = None) -> tuple[bool, str]:
    expected = _safe_csv_frame(_frame(df, columns))
    if not data.startswith(UTF8_BOM):
        return False, "Review CSV does not have a UTF-8 BOM."
    try:
        readback = pd.read_csv(io.BytesIO(data), encoding="utf-8-sig", dtype=str, keep_default_na=False)
    except Exception:
        return False, "Review CSV could not be read back."
    if list(readback.columns) != list(expected.columns) or len(readback) != len(expected):
        return False, "Review CSV schema or row count changed during export."
    for column in expected.columns:
        if readback[column].tolist() != expected[column].astype("string").fillna("").tolist():
            return False, f"Review CSV text changed in column '{column}'."
    return True, f"Review CSV verified: {len(readback)} rows; UTF-8-SIG and formula safety preserved."


def verify_review_excel_bytes(data: bytes, df: pd.DataFrame, columns: Sequence[str] | None = None) -> tuple[bool, str]:
    expected = _frame(df, columns)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        if SHEET_NAME not in wb.sheetnames:
            return False, "Review Excel sheet is missing."
        ws = wb[SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        return False, "Review Excel could not be read back."
    if not rows or list(rows[0]) != list(expected.columns) or len(rows) - 1 != len(expected):
        return False, "Review Excel schema or row count changed during export."
    for idx, column in enumerate(expected.columns):
        actual = ["" if row[idx] is None else str(row[idx]) for row in rows[1:]]
        wanted = ["" if pd.isna(v) else str(v) for v in expected[column].tolist()]
        if actual != wanted:
            return False, f"Review Excel text changed in column '{column}'."
    return True, f"Review Excel verified: {len(rows) - 1} rows; strings are preserved."


def export_review_csv(df: pd.DataFrame, product_id: str, directory: Path | None = None,
                      now: datetime | None = None, columns: Sequence[str] | None = None) -> Path:
    directory = Path(directory or EXPORT_DIR)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / make_review_filename(product_id, "csv", now)
    path.write_bytes(dataframe_to_review_csv_bytes(df, columns))
    ok, message = verify_review_csv_bytes(path.read_bytes(), df, columns)
    if not ok:
        raise ValueError(message)
    return path


def export_review_excel(df: pd.DataFrame, product_id: str, directory: Path | None = None,
                        now: datetime | None = None, columns: Sequence[str] | None = None) -> Path:
    directory = Path(directory or EXPORT_DIR)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / make_review_filename(product_id, "xlsx", now)
    path.write_bytes(dataframe_to_review_excel_bytes(df, columns))
    ok, message = verify_review_excel_bytes(path.read_bytes(), df, columns)
    if not ok:
        raise ValueError(message)
    return path
