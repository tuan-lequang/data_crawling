"""Network-free review workflow tests using synthetic inputs and a temporary cache."""
from datetime import datetime

import pandas as pd
import pytest

from config import REVIEW_COLUMNS, STANDARD_COLUMNS
from crawler.demo_reviews import DemoReviewProvider
from crawler.review_access import get_review_provider
from crawler.review_base import ReviewProviderError, ReviewRequest
from crawler.apify_reviews import Memo23ShopeeReviewProvider
from crawler.review_parser import parse_review_records
from review_pipeline import run_review_pipeline
from utils.cost_disclosure import CostModel, demo_cost_message, estimate_max_charge
from utils.csv_exporter import dataframe_to_csv_bytes, verify_csv_bytes
from utils.field_selection import select_columns, validate_field_selection
from utils.pii_redaction import redact_pii_text
from utils.review_cache import ReviewCache
from utils.review_cleaning import clean_review_dataframe, remove_duplicate_reviews
from utils.review_exporter import (
    dataframe_to_review_csv_bytes,
    dataframe_to_review_excel_bytes,
    export_review_csv,
    export_review_excel,
    verify_review_csv_bytes,
    verify_review_excel_bytes,
)
from utils.validators import validate_review_count


def request(product_id="123456", shop_id="98765", max_reviews=5, images=False):
    return ReviewRequest(product_id, shop_id, "Sản phẩm mô phỏng", max_reviews=max_reviews,
                         include_review_images=images)


def test_review_count_validation_never_accepts_zero_or_out_of_range():
    assert validate_review_count(1) == 1
    assert validate_review_count(100) == 100
    for invalid in (0, -1, True, False, "5", 101):
        with pytest.raises(ValueError):
            validate_review_count(invalid)


def test_review_parser_drops_personal_fields_and_mismatched_products():
    req = request()
    rows, mismatched = parse_review_records([
        {"reviewId": "r1", "itemId": "123456", "shopId": "98765", "ratingStar": 5,
         "comment": "Sản phẩm ổn", "createdAt": 1790000000,
         "author": "Fictional Person", "authorId": "secret-id", "authorPortrait": "private-avatar"},
        {"reviewId": "r2", "itemId": "WRONG", "shopId": "98765", "comment": "Không được nhận"},
    ], req)
    assert len(rows) == 1 and mismatched == 1
    assert rows[0]["product_id"] == "123456"
    assert set(rows[0]) == set(REVIEW_COLUMNS)
    assert not ({"author", "authorId", "authorPortrait"} & set(rows[0]))
    assert rows[0]["created_at"] == 1790000000


def test_redaction_handles_vietnamese_phone_email_url_and_contact_handle():
    source = "Gọi 0900000000, mail demo.review@example.invalid, web https://example.invalid và Zalo: fake_user_123"
    safe, count = redact_pii_text(source)
    assert count >= 4
    for secret in ("0900000000", "demo.review@example.invalid", "https://example.invalid", "fake_user_123"):
        assert secret not in safe
    assert "[phone]" in safe and "[email]" in safe and "[url]" in safe and "[contact]" in safe


def test_cleaning_redacts_before_records_can_be_cached_or_exported(tmp_path):
    req = request()
    rows, _ = parse_review_records([{
        "review_id": "r-secret", "product_id": "123456", "shop_id": "98765", "rating": 5,
        "comment": "Gọi 0912345678, mail private@example.invalid; rất ổn.",
        "created_at": "2026-10-12T12:00:00Z", "seller_reply": "Zalo: fake_contact_123",
        "data_source": "demo_synthetic", "scraped_at": "2026-10-12T12:00:00Z",
    }], req)
    df, redacted = clean_review_dataframe(rows)
    assert redacted == 2
    joined = " ".join(str(x) for x in df.astype(object).where(pd.notna(df), None).to_numpy().ravel())
    assert "0912345678" not in joined and "private@example.invalid" not in joined and "fake_contact_123" not in joined
    cache = ReviewCache(tmp_path / "cache")
    cache.put(req, "demo_synthetic", df.astype(object).where(pd.notna(df), None).to_dict("records"), 5)
    all_cache_bytes = b"".join(path.read_bytes() for path in (tmp_path / "cache").glob("*.json"))
    assert b"0912345678" not in all_cache_bytes and b"private@example.invalid" not in all_cache_bytes


def test_review_dedupe_uses_id_then_fallback_for_missing_ids():
    rows = [
        {"review_id": "same", "product_id": "1", "shop_id": "2", "rating": 5, "comment": "A"},
        {"review_id": "same", "product_id": "1", "shop_id": "2", "rating": 5, "comment": "B"},
        {"review_id": None, "product_id": "1", "shop_id": "2", "rating": 4, "comment": "No id", "created_at": "2026-01-01"},
        {"review_id": None, "product_id": "1", "shop_id": "2", "rating": 4, "comment": "No id", "created_at": "2026-01-01"},
    ]
    df, _ = clean_review_dataframe(rows)
    deduped, removed = remove_duplicate_reviews(df)
    assert removed == 2 and len(deduped) == 2


def test_demo_review_pipeline_cache_reuse_and_empty_result_ttl(tmp_path):
    cache = ReviewCache(tmp_path / "review-cache")
    req = request("123456", "98765", max_reviews=5)
    provider = DemoReviewProvider()
    first = run_review_pipeline(provider, req, cache)
    second = run_review_pipeline(provider, req, cache)
    assert not first.from_cache and second.from_cache
    assert len(first.df) == len(second.df)
    assert list(first.df.columns) == REVIEW_COLUMNS
    assert set(first.df["data_source"].dropna()) <= {"demo_synthetic"}
    # Ensure only JSON-safe, already-redacted review rows are persisted.
    for file in (tmp_path / "review-cache").glob("*.json"):
        text = file.read_text(encoding="utf-8")
        assert "0900000000" not in text and "demo.review@example.invalid" not in text
    assert cache.clear() >= 1
    assert cache.clear() == 0


def test_review_cache_does_not_reuse_entry_with_different_image_policy(tmp_path):
    cache = ReviewCache(tmp_path / "review-cache")
    no_images = request(images=False)
    cache.put(no_images, "demo_synthetic", [], 5, {"coverage": "empty"})
    with_images = request(images=True)
    assert cache.get(with_images, "demo_synthetic") is None


def test_live_review_provider_is_memo23_and_never_falls_back_to_demo():
    provider = get_review_provider("live_apify")
    assert isinstance(provider, Memo23ShopeeReviewProvider)
    with pytest.raises(ReviewProviderError, match="API token is not configured"):
        provider.fetch(request())
    assert get_review_provider("demo").data_source == "demo_synthetic"


def test_demo_reviews_are_explicitly_synthetic_and_bounded():
    provider = DemoReviewProvider()
    batch = provider.fetch(request(max_reviews=7))
    assert len(batch.records) <= 7
    assert batch.source_name == "Synthetic Demo Reviews"
    assert "not Shopee data" in batch.coverage
    assert all(r["data_source"] == "demo_synthetic" for r in batch.records)


def test_product_export_field_selection_preserves_default_schema():
    df = pd.DataFrame([{column: ("điện thoại" if column in {"product_name", "keyword"} else "1") for column in STANDARD_COLUMNS}])
    full = dataframe_to_csv_bytes(df)
    assert verify_csv_bytes(full, df)[0]
    selected = ["product_name", "price"]
    bytes_selected = dataframe_to_csv_bytes(df, selected)
    ok, message = verify_csv_bytes(bytes_selected, df, selected)
    assert ok, message
    back = pd.read_csv(pd.io.common.BytesIO(bytes_selected), encoding="utf-8-sig")
    assert list(back.columns) == selected and back.loc[0, "product_name"] == "điện thoại"
    assert len(select_columns(df, selected, STANDARD_COLUMNS).columns) == 2
    with pytest.raises(ValueError):
        validate_field_selection([], STANDARD_COLUMNS, "product output")


def sample_reviews():
    df, _ = clean_review_dataframe([
        {"review_id": "r1", "product_id": "123456", "shop_id": "98765", "rating": 5,
         "comment": "Điện thoại tốt, giao hàng nhanh", "created_at": "2026-10-12T12:00:00Z",
         "purchased_variant": "Màu đen", "like_count": 3, "seller_reply": "Cảm ơn bạn", 
         "review_image_count": 0, "review_images": None, "data_source": "demo_synthetic", "scraped_at": "2026-10-12T12:00:00Z"},
        {"review_id": "r2", "product_id": "123456", "shop_id": "98765", "rating": 4,
         "comment": "=HYPERLINK(\"https://evil.invalid\",\"click\")", "created_at": "2026-10-11T12:00:00Z",
         "purchased_variant": None, "like_count": 0, "seller_reply": None,
         "review_image_count": 0, "review_images": None, "data_source": "demo_synthetic", "scraped_at": "2026-10-12T12:00:00Z"},
    ])
    return df


def test_review_csv_export_unicode_formula_safety_and_selected_columns(tmp_path):
    df = sample_reviews()
    selected = ["review_id", "product_id", "rating", "comment"]
    data = dataframe_to_review_csv_bytes(df, selected)
    assert data.startswith(b"\xef\xbb\xbf")
    decoded = data.decode("utf-8-sig")
    assert "Điện thoại tốt" in decoded
    assert "'=HYPERLINK" in decoded
    ok, message = verify_review_csv_bytes(data, df, selected)
    assert ok, message
    path = export_review_csv(df, "123456", tmp_path, datetime(2026, 10, 13, 12, 0), selected)
    assert path.exists() and path.name.startswith("shopee_reviews_123456_")


def test_review_excel_export_uses_strings_and_roundtrips_formula_like_text(tmp_path):
    df = sample_reviews()
    data = dataframe_to_review_excel_bytes(df, ["review_id", "comment"])
    ok, message = verify_review_excel_bytes(data, df, ["review_id", "comment"])
    assert ok, message
    path = export_review_excel(df, "123456", tmp_path, datetime(2026, 10, 13, 12, 0), ["review_id", "comment"])
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, read_only=True)
    ws = wb["Reviews"]
    assert ws.cell(3, 2).data_type == "s"
    assert ws.cell(3, 2).value.startswith("=HYPERLINK")


def test_cost_model_is_provider_agnostic_and_demo_cost_message_is_clear():
    fake_model = CostModel("fictional-test-provider", per_1000_delivered=2.0, start_fee=0.01, is_synthetic=False)
    assert estimate_max_charge(5, fake_model) == 0.02
    assert "no Apify run" in demo_cost_message()
    with pytest.raises(ValueError):
        estimate_max_charge(0, fake_model)



def test_review_image_urls_are_preserved_only_when_requested():
    req = request(images=True)
    parsed, _ = parse_review_records([{
        "review_id": "img1", "product_id": "123456", "shop_id": "98765", "rating": 5,
        "comment": "Review có ảnh", "review_image_count": 1,
        "review_images": "https://demo.invalid/review-image/1.jpg", "data_source": "demo_synthetic",
    }], req)
    df, _ = clean_review_dataframe(parsed)
    assert df.loc[0, "review_images"] == "https://demo.invalid/review-image/1.jpg"


def test_review_cache_expiry_empty_negative_result_and_corrupt_file(tmp_path):
    req = request()
    expired = ReviewCache(tmp_path / "expired", ttl_seconds=0, empty_ttl_seconds=0)
    expired.put(req, "demo_synthetic", [{"review_id": "r1", "product_id": "123456", "shop_id": "98765"}], 5)
    assert expired.get(req, "demo_synthetic") is None
    expired.put(req, "demo_synthetic", [], 5)
    assert expired.get(req, "demo_synthetic") is None

    cache = ReviewCache(tmp_path / "corrupt")
    cache.directory.mkdir(parents=True)
    corrupt_path = cache.directory / f"{cache._key(req, 'demo_synthetic')}.json"
    corrupt_path.write_text("{ definitely not json", encoding="utf-8")
    assert cache.get(req, "demo_synthetic") is None
    assert not corrupt_path.exists()


def test_review_cache_can_reuse_underfilled_result_for_larger_request(tmp_path):
    cache = ReviewCache(tmp_path / "covering")
    small = request(max_reviews=10)
    cached_records = [{"review_id": f"r{i}", "product_id": "123456", "shop_id": "98765"} for i in range(3)]
    cache.put(small, "demo_synthetic", cached_records, 10)
    larger = request(max_reviews=50)
    entry = cache.get(larger, "demo_synthetic")
    assert entry is not None and len(entry.records) == 3
