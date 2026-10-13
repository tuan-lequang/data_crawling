import re
from datetime import datetime

import pandas as pd

from config import STANDARD_COLUMNS
from utils.cleaning import clean_dataframe
from utils.csv_exporter import (
    UTF8_BOM,
    dataframe_to_csv_bytes,
    dataframe_to_excel_bytes,
    export_csv,
    export_excel,
    make_excel_filename,
    make_filename,
    verify_csv_bytes,
    verify_excel_bytes,
)


def sample_df():
    return clean_dataframe(pd.DataFrame([
        {"product_id": 22093519050, "product_name": "Giày Sneaker Nam Đế Êm, \"Hot\"", "shop_id": 5,
         "shop_name": "Cửa hàng giày Việt", "shop_location": "Thành phố Hồ Chí Minh",
         "category": "Điện thoại & phụ kiện", "keyword": "điện thoại", "price": 219000,
         "rating": 4.8, "sold": 1500, "is_shopee_verified": True, "data_source": "demo"},
        {"product_id": 2, "product_name": "Áo thun nam", "sold": None, "rating": None},
    ]))


def test_filename_format():
    assert make_filename(datetime(2026, 10, 7, 13, 55, 0)) == "shopee_products_20261007_135500.csv"
    assert re.fullmatch(r"shopee_products_\d{8}_\d{6}\.csv", make_filename())


def test_bytes_are_utf8_sig():                                    # test 12
    data = dataframe_to_csv_bytes(sample_df())
    assert data.startswith(UTF8_BOM)
    assert "Giày Sneaker Nam Đế Êm" in data.decode("utf-8-sig")


def test_all_20_columns_in_order():
    back = pd.read_csv(pd.io.common.BytesIO(dataframe_to_csv_bytes(sample_df())), encoding="utf-8-sig", dtype=str)
    assert list(back.columns) == STANDARD_COLUMNS and len(STANDARD_COLUMNS) == 20


def test_roundtrip_values_and_empty_cells():
    df = sample_df()
    ok, message = verify_csv_bytes(dataframe_to_csv_bytes(df), df)
    assert ok, message
    back = pd.read_csv(pd.io.common.BytesIO(dataframe_to_csv_bytes(df)), encoding="utf-8-sig",
                       dtype=str, keep_default_na=False)
    assert back.loc[0, "rating"] == "4.8" and back.loc[0, "price"] == "219000"
    assert back.loc[1, "sold"] == "" and back.loc[1, "rating"] == ""        # nulls stay empty
    assert back.loc[0, "product_id"] == "22093519050"                       # ids are not mangled


def test_verify_detects_missing_bom_and_changed_text():
    df = sample_df()
    data = dataframe_to_csv_bytes(df)
    assert verify_csv_bytes(data[len(UTF8_BOM):], df)[0] is False           # plain utf-8, no BOM
    other = df.copy()
    other.loc[0, "product_name"] = "Khác"
    assert verify_csv_bytes(data, other)[0] is False
    other = df.copy()
    other.loc[0, "keyword"] = "dien thoai"
    assert verify_csv_bytes(data, other)[0] is False                        # check non-title text too


def test_export_csv_writes_a_readable_file(tmp_path):
    df = sample_df()
    path = export_csv(df, tmp_path, now=datetime(2026, 10, 7, 21, 5, 0))
    assert path.name == "shopee_products_20261007_210500.csv"
    assert path.read_bytes().startswith(UTF8_BOM)
    back = pd.read_csv(path, encoding="utf-8-sig")
    assert back.loc[0, "product_name"] == 'Giày Sneaker Nam Đế Êm, "Hot"'
    assert back.loc[0, "shop_location"] == "Thành phố Hồ Chí Minh"
    assert back.loc[0, "keyword"] == "điện thoại"


def test_empty_dataframe_exports_header_only():
    data = dataframe_to_csv_bytes(clean_dataframe(pd.DataFrame(columns=STANDARD_COLUMNS)))
    assert data.decode("utf-8-sig").strip().split(",") == STANDARD_COLUMNS


def test_excel_filename_format():
    assert make_excel_filename(datetime(2026, 10, 7, 13, 55, 0)) == "shopee_products_20261007_135500.xlsx"
    assert re.fullmatch(r"shopee_products_\d{8}_\d{6}\.xlsx", make_excel_filename())


def test_excel_roundtrip_and_schema():
    df = sample_df()
    data = dataframe_to_excel_bytes(df)
    ok, message = verify_excel_bytes(data, df)
    assert ok, message
    from openpyxl import load_workbook
    wb = load_workbook(pd.io.common.BytesIO(data), read_only=True, data_only=True)
    ws = wb["Products"]
    headers = [cell.value for cell in ws[1]]
    assert headers == STANDARD_COLUMNS and ws.max_row == len(df) + 1
    assert ws[2][0].value == "22093519050"  # product_id remains text, not scientific notation
    assert ws[2][1].value == 'Giày Sneaker Nam Đế Êm, "Hot"'
    assert ws.cell(row=2, column=headers.index("keyword") + 1).value == "điện thoại"


def test_export_excel_writes_readable_file(tmp_path):
    df = sample_df()
    path = export_excel(df, tmp_path, now=datetime(2026, 10, 7, 21, 5, 0))
    assert path.name == "shopee_products_20261007_210500.xlsx"
    ok, message = verify_excel_bytes(path.read_bytes(), df)
    assert ok, message
